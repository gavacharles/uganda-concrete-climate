"""Build submission-style Word documents from the Claude Docs markdown exports.

  .venv/bin/python src/build_docx.py A manuscripts/paperA_export.md
  .venv/bin/python src/build_docx.py B manuscripts/paperB_export.md

The working-notes section is dropped, image placeholders are replaced by the local high-resolution
figures, LaTeX blocks become Word equations, and the title and author block are set for submission.
Output: manuscripts/PaperA_manuscript.docx or manuscripts/PaperB_manuscript.docx
"""
import re
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"

PAPERS = {
    "A": dict(
        title="Climate change and carbonation-induced corrosion risk of a rapidly growing concrete building stock: "
              "a probabilistic, building-age-aware assessment for Uganda",
        out="PaperA_manuscript.docx",
        images={
            "Figure 1. Assessment framework": "fig1_framework.png",
            "Figure 2. Projected warming": "fig2_climate.png",
            "Figure 3. Corrosion-initiation probability": "fig3_site_pinit.png",
            "Figure 4. Construction quality": "fig4_quality_vs_climate.png",
            "Figure 5. National building-stock": "fig5_stock_trajectory.png",
            "Figure 6. Map of stock exposure": "fig6_stock_map.png",
            "Figure 7. Extra cover": "fig7_adaptation.png",
            "Figure 8. Sensitivity": "fig8_sensitivity.png",
            "Figure S0. PRISMA": "prisma_flow_scopus.png",
            "Figure 9. Probability of cover cracking": "fig9_cracking.png",
            "Figure 10. New construction": "fig10_new_construction.png",
            "Figure S3. Measured versus fib": "figS3_carbonation_check.png",
        }),
    "B": dict(
        title="Neighbourhood data, urban heat and heavier rain: exposure of a fast-growing tropical building stock "
              "to climate-driven concrete deterioration in Uganda",
        out="PaperB_manuscript.docx",
        images={
            "Figure 1. Growth": "fig11_growth.png",
            "Figure 2. Machine-learning": "fig13_machine_learning.png",
            "Figure 3. Urban heat": "fig12_urban_heat.png",
            "Figure 4. Rainfall": "fig14_rain_extremes.png",
        }),
}
AUTHOR = "Gavamukulya Charles"
AFFILIATION = "[Affiliation, address, ORCID and corresponding-author e-mail to be added]"


def preprocess(md, cfg):
    # drop the title line, byline and working notes (everything before the Abstract heading)
    md = md[md.index("## Abstract"):]
    md = md.replace("&#91;", "[")

    def image(m):
        alt = m.group(1).strip()
        for key, f in cfg["images"].items():
            if alt.startswith(key):
                return f"![]({(FIG / f).as_posix()}){{width=16cm}}"
        raise KeyError(f"no figure file for placeholder: {alt}")

    md = re.sub(r"\[(?:image|embedded content): ([^\]]*?)\\?\]", image, md)
    md = re.sub(r"```latex\n(.*?)\n```", lambda m: "$$" + m.group(1).strip() + "$$", md, flags=re.S)
    return md


def finish(path, title):
    doc = Document(path)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for p in doc.paragraphs:
        if p.style.name in ("Body Text", "First Paragraph", "Normal", "Compact"):
            p.paragraph_format.line_spacing = 1.5
    body = doc.element.body
    # title and author block at the top
    first = doc.paragraphs[0]
    t = first.insert_paragraph_before(title)
    t.style = doc.styles["Title"]
    a = first.insert_paragraph_before(AUTHOR)
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f = first.insert_paragraph_before(AFFILIATION)
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # continuous line numbering, as most journals ask for at review
    for sect in doc.sections:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sect._sectPr.append(ln)
        # page numbers in the footer
        para = sect.footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
    _ = body
    doc.save(path)


def main(which, src):
    cfg = PAPERS[which]
    md = preprocess(Path(src).read_text(), cfg)
    out = ROOT / "manuscripts" / cfg["out"]
    pypandoc.convert_text(md, "docx", format="markdown+tex_math_dollars+pipe_tables", outputfile=str(out),
                          extra_args=["--resource-path", str(ROOT)])
    finish(out, cfg["title"])
    print("wrote", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
